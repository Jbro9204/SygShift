from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from pathlib import Path

from docx import Document
from pypdf import PdfReader


def table_lines(path: Path) -> list[str]:
    document = Document(path)
    lines = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            value = " | ".join(cell.text.strip().replace("\n", " / ") for cell in row.cells)
            if value.strip(" |"):
                lines.append(value)
    return lines


def after_label(lines: list[str], label: str, fallback: str = "") -> str:
    prefix = label.upper() + ":"
    for line in lines:
        if line.upper().startswith(prefix):
            return line[len(prefix) :].strip()
    return fallback


def record_class(lines: list[str], fallback: str) -> str:
    for line in lines[:15]:
        match = re.search(r"RECORD CLASS\s*\|\s*(.+?)(?:\s*\|\s*[A-Z][A-Z /&-]+\s*\||$)", line)
        if match:
            return match.group(1).strip()
    return fallback


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def pdf_metadata(path: Path) -> dict[str, object]:
    reader = PdfReader(path)
    text = "\n".join((page.extract_text() or "") for page in reader.pages)
    return {
        "pageCount": len(reader.pages),
        "sizeBytes": path.stat().st_size,
        "sha256": sha256(path),
        "searchText": re.sub(r"\s+", " ", text).strip(),
    }


def sensitivity_for(section: str) -> tuple[str, str, str]:
    if "Leave_CONF" in section:
        return "hr-medical", "highly_restricted", "hr_only"
    if "ER_RESTRICT" in section or "/05_Performance" in section:
        return "hr-disciplinary", "highly_restricted", "hr_only"
    if any(token in section for token in ("Safety", "Security_Ops", "Safety_Fleet", "Separation", "Records_Legal")):
        return "hr-legal-safety", "restricted", "hr_only"
    if "Pay_Benefits" in section:
        return "hr-financial", "highly_restricted", "hr_only"
    return "hr-general", "confidential", "hr_only"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-root", required=True, type=Path)
    parser.add_argument("--pdf-root", required=True, type=Path)
    parser.add_argument("--output-root", required=True, type=Path)
    args = parser.parse_args()

    source_root = args.source_root.resolve()
    pdf_root = args.pdf_root.resolve()
    output_root = args.output_root.resolve()
    output_root.mkdir(parents=True, exist_ok=True)

    with (source_root / "HR/00_README/FILE_MAP.csv").open(encoding="utf-8-sig", newline="") as stream:
        hr_rows = list(csv.DictReader(stream))
    with (source_root / "TRN/00_START/CROSSWALK.csv").open(encoding="utf-8-sig", newline="") as stream:
        crosswalk = {row["source_safe_path"]: row for row in csv.DictReader(stream)}
    with (source_root / "TRN/00_START/PROGRAM_REGISTER.csv").open(encoding="utf-8-sig", newline="") as stream:
        training_rows = list(csv.DictReader(stream))
    with (source_root / "TRN/00_START/MODULE_REGISTER.csv").open(encoding="utf-8-sig", newline="") as stream:
        module_rows = {row["code"]: row for row in csv.DictReader(stream)}

    items: list[dict[str, object]] = []
    codes: set[str] = set()
    for row in hr_rows:
        safe_path = row["safe_path"].replace("/", "\\")
        if safe_path == "01_FORMS\\01_Recruit\\GS-HR-101.docx":
            continue
        code = row["code"]
        title = row["title"]
        if safe_path == "03_INTERVIEW\\02_Core\\GS-HR-101_v2.docx":
            code = "GS-HR-101"
            title = "Candidate Interview Evaluation v2.0"
        if code in codes:
            raise ValueError(f"Duplicate canonical HR code: {code}")
        codes.add(code)
        source_path = source_root / "HR" / safe_path
        pdf_path = pdf_root / "HR" / Path(safe_path).with_suffix(".pdf")
        if not pdf_path.exists():
            raise FileNotFoundError(pdf_path)
        lines = table_lines(source_path)
        vault, sensitivity, audience = sensitivity_for(row["section"])
        guide = crosswalk.get(row["safe_path"])
        items.append({
            "kind": "hr_source",
            "code": code,
            "title": title,
            "category": row["category"],
            "section": row["section"],
            "recordClass": record_class(lines, row["category"]),
            "purpose": after_label(lines, "PURPOSE", title),
            "audience": audience,
            "sensitivity": sensitivity,
            "vaultCode": vault,
            "status": "draft_for_adoption",
            "sourceRelativePath": str(source_path.relative_to(source_root)).replace("\\", "/"),
            "pdfRelativePath": str(pdf_path.relative_to(output_root)).replace("\\", "/"),
            "relatedModules": (guide or {}).get("related_modules", "").split("; ") if guide else [],
            "guideCode": (guide or {}).get("d_guide"),
            **pdf_metadata(pdf_path),
        })

    for row in training_rows:
        relative = row["path"].replace("/", "\\")
        source_path = source_root / "TRN" / relative
        pdf_path = pdf_root / "TRN" / Path(relative).with_suffix(".pdf")
        if not pdf_path.exists():
            raise FileNotFoundError(pdf_path)
        series = row["series"]
        kind = {"P": "training_admin", "T": "training_module", "D": "document_guide", "F": "training_form"}[series]
        module = module_rows.get(row["code"])
        lines = table_lines(source_path)
        training_record_class = {
            "P": "Training Program Administration",
            "T": "Training Course Material",
            "D": "Controlled Document Use Guide",
            "F": "Training Administration Record",
        }[series]
        related_modules = [row["code"]] if series == "T" else []
        if series == "D":
            crosswalk_row = next((value for value in crosswalk.values() if value.get("d_guide") == row["code"]), None)
            related_modules = crosswalk_row.get("related_modules", "").split("; ") if crosswalk_row else []
        items.append({
            "kind": kind,
            "code": row["code"],
            "title": row["title"],
            "category": row["category"],
            "section": f"TRN/{Path(relative).parent.as_posix()}",
            "recordClass": record_class(lines, training_record_class),
            "purpose": after_label(lines, "PURPOSE", row["title"]),
            "audience": "hr_only",
            "sensitivity": "confidential",
            "vaultCode": "hr-general",
            "status": "draft_for_adoption",
            "sourceRelativePath": str(source_path.relative_to(source_root)).replace("\\", "/"),
            "pdfRelativePath": str(pdf_path.relative_to(output_root)).replace("\\", "/"),
            "relatedModules": related_modules,
            "guideCode": row["code"] if series == "D" else None,
            "audienceDescription": module["audience"] if module else None,
            "expectedMinutes": int(module["minutes"]) if module else None,
            "passingStandard": module["passing_standard"] if module else None,
            "sourceCodes": module["source_codes"].split("; ") if module else [],
            **pdf_metadata(pdf_path),
        })

    expected = 537
    if len(items) != expected:
        raise ValueError(f"Expected {expected} canonical PDFs, found {len(items)}")
    if len({item["sha256"] for item in items}) != expected:
        raise ValueError("Duplicate canonical PDF content detected")

    manifest = {
        "package": "Guardianship Security HR System v2.1 with Training",
        "libraryVersion": "2.1",
        "status": "draft_for_company_adoption",
        "canonicalPdfCount": len(items),
        "excluded": [{
            "path": "HR/01_FORMS/01_Recruit/GS-HR-101.docx",
            "reason": "Superseded by the materially more complete GS-HR-101 v2.0 interview evaluation.",
        }],
        "items": items,
    }
    (output_root / "catalog.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    print(f"Validated {len(items)} canonical PDFs and wrote {output_root / 'catalog.json'}")


if __name__ == "__main__":
    main()
